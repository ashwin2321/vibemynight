import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"C:\Users\ADMIN\Downloads\vibemynight-phase1-10\documentation"
os.makedirs(OUT_DIR, exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    if level == 1:
        run.font.name = 'Segoe UI'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138) # Deep Blue
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.name = 'Segoe UI'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(131, 24, 67) # Berry/Pink
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.name = 'Segoe UI'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(55, 65, 81) # Slate
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
    return h

def add_callout(doc, text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border thick accent
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="3B82F6"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor(30, 58, 138)
    
    r_text = p.add_run(text)
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(31, 41, 55)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_table_data(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    # Header
    hdr_row = tbl.rows[0]
    for idx, h in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        if col_widths and idx < len(col_widths):
            cell.width = col_widths[idx]
            
    # Data rows
    for r_idx, row_data in enumerate(data):
        row = tbl.rows[r_idx + 1]
        bg = "F9FAFB" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.name = "Segoe UI"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(17, 24, 39)
            if col_widths and c_idx < len(col_widths):
                cell.width = col_widths[c_idx]
                
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# BUILD SRS DOCUMENT
doc = Document()

# Cover / Header
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(36)
title_p.paragraph_format.space_after = Pt(6)
r_title = title_p.add_run("SOFTWARE REQUIREMENTS SPECIFICATION (SRS)")
r_title.font.name = "Segoe UI"
r_title.font.size = Pt(22)
r_title.font.bold = True
r_title.font.color.rgb = RGBColor(30, 58, 138)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_after = Pt(24)
r_sub = sub_p.add_run("VibeMyNight — Event Discovery & Pass Inquiry Platform (v1.0)")
r_sub.font.name = "Segoe UI"
r_sub.font.size = Pt(13)
r_sub.font.color.rgb = RGBColor(131, 24, 67)

meta_data = [
    ["Document Version", "1.0.0 (Production Release)"],
    ["Author / Architecture", "VibeMyNight Core Engineering Team"],
    ["Target Platform", "Web (Vercel Edge), Android, iOS"],
    ["Backend Stack", "Spring Boot 3.3.4, Java 21, MySQL 8, JPA/Hibernate, JWT"],
    ["Frontend Stack", "Flutter 3.27+ (Dart), Riverpod, GoRouter, Dark Neon Design"],
    ["Status", "Approved & Feature Complete (Phase 1–14)"]
]
add_table_data(doc, ["Attribute", "Specification Details"], meta_data, [Inches(2.2), Inches(4.3)])

doc.add_page_break()

# 1. Executive Overview & Business Model
add_styled_heading(doc, "1. Project Overview & Business Model", 1)
p = doc.add_paragraph()
p.add_run("VibeMyNight is a premier digital nightlife and cultural event discovery platform, specifically optimized for high-volume Indian seasonal festivals (such as Navratri Mahotsav, Sunburn Arenas, New Year Galas, and Concerts). ")
p.add_run("The platform connects event organizers with thousands of festival attendees seeking passes, VIP tables, group bookings, and daily artist schedules.\n\n")

add_styled_heading(doc, "1.1 Core Business Model (v1.0 WhatsApp Lead Routing)", 2)
p2 = doc.add_paragraph()
p2.add_run("Traditional event ticketing platforms charge prohibitive 8–15% convenience fees per ticket and force instant automated payment gateway checkouts that often lead to cart abandonment and high payment failure rates during peak flash sales. ")
p2.add_run("VibeMyNight v1.0 implements a frictionless WhatsApp-driven Pass Inquiry Engine:\n")

bullets = [
    ("Dynamic Catalog Exploration: ", "Customers browse high-resolution festival hero cards, daily performing artists, venue amenities, rules, and multiple pass tiers without requiring account registration."),
    ("Server-Verified Pass Inquiries: ", "Customers select a festival date and pass category (e.g. VIP Couple, Ground General). The server validates rates from the database, records an immutable inquiry record, and generates a pre-formatted WhatsApp deep link."),
    ("WhatsApp High-Touch Conversion: ", "The customer is seamlessly redirected to WhatsApp (`https://wa.me/917041615131?text=...`) with their unique inquiry number, event date, pass category, quantity, and estimated total."),
    ("Zero Commission / Direct Booking: ", "Organizers and admin staff close pass bookings directly via WhatsApp Pay, UPI, cash, or direct VIP bank transfers, completely avoiding intermediary gateway fees.")
]
for b_title, b_desc in bullets:
    bp = doc.add_paragraph(style='List Bullet')
    r_bt = bp.add_run(b_title)
    r_bt.bold = True
    r_bt.font.color.rgb = RGBColor(30, 58, 138)
    bp.add_run(b_desc)

add_callout(doc, "Payment Gateways (Razorpay/Stripe) are deliberately omitted in v1.0 to maximize conversion speed, eliminate transaction fees, and enable personalized high-touch VIP table allocation via direct WhatsApp chat.", "BUSINESS ARCHITECTURE")

# 2. User Roles & Personas
add_styled_heading(doc, "2. User Roles & Personas", 1)
roles_data = [
    ["Role", "Access Channel", "Authentication", "Permissions & Capabilities"],
    ["Guest / Customer", "Public Web / Mobile App", "None required (Zero Login Friction)", "Browse festival catalog, view multi-day schedules, inspect artist portfolios, submit pass inquiries, open direct WhatsApp booking."],
    ["Platform Admin", "Admin Portal (/admin/login)", "JWT Bearer Token (BCrypt Auth)", "Full CRUD on Events, Festival Days, Pass Tiers, Artists, Facilities, Platform Settings; manage inquiry leads ledger; generate offline bills & WhatsApp receipts."]
]
add_table_data(doc, roles_data[0], roles_data[1:], [Inches(1.5), Inches(1.6), Inches(1.8), Inches(2.2)])

# 3. Functional Requirements
add_styled_heading(doc, "3. Functional Requirements", 1)

add_styled_heading(doc, "3.1 Module 1: Event & Festival Catalog (FR-EVT)", 2)
evt_reqs = [
    ["FR-EVT-01", "Multi-Day Structure", "The system shall support multi-day festivals (e.g., 9 nights of Navratri) where an overarching event entity contains 1 to N calendar performance days."],
    ["FR-EVT-02", "Event Lifecycle States", "Events shall transition through lifecycle states: DRAFT (editable, hidden from public), PUBLISHED (live on customer catalog), CANCELLED, and ARCHIVED."],
    ["FR-EVT-03", "Rich Media & Guidelines", "Events shall store and render hero cover banners, thumbnails, photo gallery carousels, highlights, and entry guidelines/rules."],
    ["FR-EVT-04", "SEO Friendly Slugs", "Each event shall have a unique URL slug (e.g. /events/navratri-nights-2026) for deep linking and social media sharing."]
]
add_table_data(doc, ["Req ID", "Feature", "Functional Specification"], evt_reqs, [Inches(1.2), Inches(1.8), Inches(4.0)])

add_styled_heading(doc, "3.2 Module 2: Multi-Day Lineup & Differential Pricing (FR-DAY & FR-PASS)", 2)
day_reqs = [
    ["FR-DAY-01", "Calendar Day Scheduling", "Each festival day shall maintain its date, day number index (Day 1..N), start/end curfew timings, and special program title."],
    ["FR-DAY-02", "Artist Lineup Assignment", "Admin can assign 1 to N artists to each day, designating exactly 1 Primary Headliner and setting individual performance slots."],
    ["FR-PASS-01", "Tiered Pass Categories", "Each festival day can have multiple independent ticket tiers (VIP, Couple, Female Stag, Male Stag, Group of 5)."],
    ["FR-PASS-02", "Differential Dynamic Pricing", "Prices vary per day (e.g. Day 1: ₹499, Weekend/Sharad Purnima: ₹1499). The system computes starting prices dynamically."]
]
add_table_data(doc, ["Req ID", "Feature", "Functional Specification"], day_reqs, [Inches(1.2), Inches(1.8), Inches(4.0)])

add_styled_heading(doc, "3.3 Module 3: Pass Inquiry & WhatsApp Routing Engine (FR-INQ)", 2)
inq_reqs = [
    ["FR-INQ-01", "Anti-Tamper Price Compute", "When a customer submits an inquiry, the server retrieves the true unit price from MySQL, ignoring any client-sent price values."],
    ["FR-INQ-02", "Unique Inquiry Reference", "The backend generates a collision-free human-readable reference code (e.g., VMN-2026-8821)."],
    ["FR-INQ-03", "WhatsApp Deep Link Dispatch", "The system constructs a formatted URI (https://wa.me/<number>?text=...) with inquiry summary, event name, pass category, quantity, and total."],
    ["FR-INQ-04", "Inquiry Status Ledger", "Inquiries are stored in DB with lifecycle: NEW -> CONTACTED -> CONFIRMED -> CANCELLED."]
]
add_table_data(doc, ["Req ID", "Feature", "Functional Specification"], inq_reqs, [Inches(1.2), Inches(1.8), Inches(4.0)])

add_styled_heading(doc, "3.4 Module 4: Admin Billing & Invoice Generator (FR-BILL)", 2)
p_bill = doc.add_paragraph()
p_bill.add_run("The Admin Billing Module (/admin/billing) enables manual or inquiry-backed bill creation for offline pass distribution:\n")
b_items = [
    ("Inquiry Auto-Populate: ", "Selecting an existing inquiry auto-fills customer details, pass rates, and event info in 1-click."),
    ("Manual Entry Mode: ", "Allows generating bills for walk-in or phone inquiries without requiring existing backend records."),
    ("Live Financial Computation: ", "Real-time calculation of Subtotal, Discount, Tax/GST %, and Grand Total."),
    ("WhatsApp Invoice Sharing: ", "Generates formatted WhatsApp text receipts and opens direct chat with customer (`wa.me/<phone>?text=...`).")
]
for bt, bd in b_items:
    bp = doc.add_paragraph(style='List Bullet')
    r_bt = bp.add_run(bt)
    r_bt.bold = True
    bp.add_run(bd)

# 4. Non-Functional Requirements
add_styled_heading(doc, "4. Non-Functional Requirements (NFR)", 1)
nfr_data = [
    ["NFR ID", "Quality Attribute", "Requirement & Metric Benchmark"],
    ["NFR-SEC-01", "JWT RBAC Security", "All /admin/** REST endpoints require valid HMAC-SHA256 JWT tokens with ROLE_ADMIN authority. Passwords hashed using BCrypt (12 rounds)."],
    ["NFR-PERF-01", "Response Time", "Public catalog APIs shall respond within < 150ms under 500 concurrent connections."],
    ["NFR-RESP-01", "Cross-Platform UI", "Flutter Web, Android, and iOS shall render 100% responsive dark-neon layouts matching Figma specifications with zero layout overflow errors."],
    ["NFR-AVAIL-01", "High Availability", "Frontend deployed on Vercel Global Edge CDN (99.99% uptime); Backend configured for containerized cloud scaling (Docker)."],
    ["NFR-DATA-01", "Data Integrity", "Server-side price computation and database foreign key cascades prevent dangling ticket categories or orphaned inquiries."]
]
add_table_data(doc, nfr_data[0], nfr_data[1:], [Inches(1.2), Inches(1.8), Inches(4.0)])

# Output
file_path = os.path.join(OUT_DIR, "1_Software_Requirements_Specification_VibeMyNight.docx")
doc.save(file_path)
print("SUCCESS: 1_Software_Requirements_Specification_VibeMyNight.docx generated at", file_path)
