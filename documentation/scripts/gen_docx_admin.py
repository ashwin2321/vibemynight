import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"C:\Users\ADMIN\Downloads\vibemynight-phase1-10\documentation"
os.makedirs(OUT_DIR, exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    if level == 1:
        run.font.name = 'Segoe UI'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.name = 'Segoe UI'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(131, 24, 67)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.name = 'Segoe UI'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(55, 65, 81)
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
    return h

def add_callout(doc, text, title="ADMIN PRO-TIP"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="831843"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor(131, 24, 67)
    r_text = p.add_run(text)
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(31, 41, 55)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_table_data(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    hdr_row = tbl.rows[0]
    for idx, h in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "831843") # Berry / Admin Theme
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        if col_widths and idx < len(col_widths):
            cell.width = col_widths[idx]
            
    for r_idx, row_data in enumerate(data):
        row = tbl.rows[r_idx + 1]
        bg = "F9FAFB" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.name = "Segoe UI"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(17, 24, 39)
            if col_widths and c_idx < len(col_widths):
                cell.width = col_widths[c_idx]
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# BUILD ADMIN MANUAL
doc = Document()

# Cover
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(36)
title_p.paragraph_format.space_after = Pt(6)
r_title = title_p.add_run("USER MANUAL — ADMIN MANAGEMENT PORTAL")
r_title.font.name = "Segoe UI"
r_title.font.size = Pt(22)
r_title.font.bold = True
r_title.font.color.rgb = RGBColor(131, 24, 67)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_after = Pt(24)
r_sub = sub_p.add_run("VibeMyNight — Festival CMS, Pass Pricing, Lead Management & Billing Generator")
r_sub.font.name = "Segoe UI"
r_sub.font.size = Pt(13)
r_sub.font.color.rgb = RGBColor(30, 58, 138)

meta_data = [
    ["Portal Version", "1.0.0 (Production)"],
    ["Admin URL", "https://vibemynight.vercel.app/admin/login"],
    ["Authentication", "JWT Token with BCrypt Encrypted Storage"],
    ["Default Admin Credentials", "Username: admin | Password: (Configured via DB seed)"],
    ["Key Modules", "Dashboard, Events, Multi-Day, Passes, Artists, Facilities, Inquiries, Billing, Settings"]
]
add_table_data(doc, ["Portal Attribute", "Specification Details"], meta_data, [Inches(2.2), Inches(4.3)])

doc.add_page_break()

# 1. Admin Authentication & Navigation
add_styled_heading(doc, "1. Admin Login & Dashboard Overview", 1)
p_auth = doc.add_paragraph()
p_auth.add_run("The Admin Portal is secured via JWT token authentication. Unauthenticated access attempts to any `/admin/**` routes are automatically intercepted and redirected to `/admin/login`.\n")

add_styled_heading(doc, "1.1 Dashboard Overview (Route: /admin/dashboard)", 2)
p_dash = doc.add_paragraph()
p_dash.add_run("The Admin Dashboard provides real-time operational visibility:")
d_metrics = [
    ("Total Inquiries Card: ", "Live count of all pass booking leads received across all festivals."),
    ("New Leads Metric: ", "Highlights pending customer inquiries awaiting organizer response."),
    ("Estimated Pipeline Revenue: ", "Sum total of INR value represented by incoming pass inquiries."),
    ("Quick Action Toolbar: ", "1-click shortcuts for '+ Create Event', '+ Add Artist', 'Excel / CSV Events', 'Generate Bill / WhatsApp', and 'Manage Inquiries'."),
    ("Recent Inquiries Feed: ", "Live table of the latest 6 customer leads with direct status change actions.")
]
for dm, dd in d_metrics:
    bp = doc.add_paragraph(style='List Bullet')
    r_dm = bp.add_run(dm)
    r_dm.bold = True
    bp.add_run(dd)

# 2. Festival & Multi-Day Management
add_styled_heading(doc, "2. Event, Day & Pass Management", 1)

add_styled_heading(doc, "2.1 Managing Events (Route: /admin/events)", 2)
p_evtm = doc.add_paragraph()
p_evtm.add_run("1. Navigate to 'Events' in sidebar.\n")
p_evtm.add_run("2. Click '+ Create Event' to launch the Event Editor.\n")
p_evtm.add_run("3. Enter Festival Title, SEO Slug, City, Venue, Start/End dates, and Cover Banner URL.\n")
p_evtm.add_run("4. Add Photo Gallery image URLs, Highlights, and Entry Guidelines.\n")
p_evtm.add_run("5. Set status to DRAFT while setting up days/passes, and switch to PUBLISHED when ready for public sales.")

add_styled_heading(doc, "2.2 Festival Days & Artist Lineup (/admin/events/:id/days)", 2)
p_daym = doc.add_paragraph()
p_daym.add_run("Each festival can have multiple independent performance days (e.g. Navratri Day 1 to Day 9):\n")
p_daym.add_run("- Click 'Manage Days' on any event.\n")
p_daym.add_run("- Add a day with its calendar date and special program title.\n")
p_daym.add_run("- Under 'Lineup', assign registered artists, mark the Primary Headliner, and set performance timing.")

add_styled_heading(doc, "2.3 Pass Categories & Live Pricing (/admin/event-days/:dayId/passes)", 2)
p_passm = doc.add_paragraph()
p_passm.add_run("To configure pass tiers for a specific day:\n")
p_passm.add_run("1. Click 'Passes' on the target festival day.\n")
p_passm.add_run("2. Click '+ Add Pass Category'.\n")
p_passm.add_run("3. Specify Tier Name (VIP, Couple, Student, General), Price in INR, and Allocation Capacity.\n")
p_passm.add_run("4. Prices update dynamically on public customer date pickers.")

# 3. Master Catalogs (Artists & Facilities)
add_styled_heading(doc, "3. Master Catalogs (Artists & Facilities)", 1)
p_art = doc.add_paragraph()
p_art.add_run("Artists and Facilities are managed in centralized master catalogs for cross-event reuse:\n")
p_art.add_run("- **Artists CMS (/admin/artists):** Add performer names, genre tags, bio, portrait photo URL, Spotify link, and Instagram URL.\n")
p_art.add_run("- **Facilities CMS (/admin/facilities):** Add reusable venue amenities (Valet Parking, Food Court, VIP Lounge) with Material Icons.")

# 4. Inquiry Lead Management
add_styled_heading(doc, "4. Customer Inquiries Ledger (/admin/inquiries)", 1)
p_inqm = doc.add_paragraph()
p_inqm.add_run("The Inquiries Ledger records all customer pass booking leads:\n")
p_inqm.add_run("1. Filter leads by status (NEW, CONTACTED, CONFIRMED, CANCELLED) or search by customer name/mobile.\n")
p_inqm.add_run("2. Click on an inquiry to view full lead breakdown: pass tier requested, quantity, calculated total, and customer notes.\n")
p_inqm.add_run("3. Update status from NEW -> CONTACTED -> CONFIRMED as offline payment is received.")

# 5. Admin Billing & WhatsApp Invoice Generator
add_styled_heading(doc, "5. Billing & WhatsApp Invoice Generator (/admin/billing)", 1)
p_bill = doc.add_paragraph()
p_bill.add_run("The Billing Module empowers organizers to generate, calculate, and send instant receipts to guests:")

bill_features = [
    ("Dual Creation Modes: ", "1) Auto-Populate from Backend Inquiry: Select any inquiry from dropdown to populate customer and pass data instantly. 2) 100% Manual Entry: Enter walk-in customer details freely without existing backend records."),
    ("Live Financial Breakdown: ", "Adjust Price, Quantity (+/- stepper), Discount (₹), Tax / GST %, and choose Payment Method (UPI, Cash, Card, WhatsApp Pay)."),
    ("Dark-Neon Receipt Preview: ", "Renders an elegant digital invoice matching Figma brand guidelines with VibeMyNight branding and status badge."),
    ("Direct WhatsApp Sharing: ", "Click 'Send on WhatsApp' to generate a formatted text receipt and launch WhatsApp (`wa.me/<phone>?text=...`) directly to the customer."),
    ("One-Tap Clipboard Copy: ", "Click 'Copy Formatted Bill' to paste invoice text into email, SMS, or CRM systems.")
]
for bf, bd in bill_features:
    bp = doc.add_paragraph(style='List Bullet')
    r_bf = bp.add_run(bf)
    r_bf.bold = True
    r_bf.font.color.rgb = RGBColor(131, 24, 67)
    bp.add_run(bd)

add_callout(doc, "The Billing Module is completely offline-resilient — bills can be calculated, formatted, and shared over WhatsApp even when no backend database connection is present.", "OPERATIONAL RESILIENCE")

# 6. Platform Settings
add_styled_heading(doc, "6. Platform Configuration (/admin/settings)", 1)
p_set = doc.add_paragraph()
p_set.add_run("Manage global platform identity and customer care contact points:\n")
p_set.add_run("- **Brand Name:** Displayed on headers, receipts, and meta tags.\n")
p_set.add_run("- **Official WhatsApp Helpline Number:** The primary WhatsApp contact that receives all customer pass inquiries (e.g. `917041615131`).\n")
p_set.add_run("- **Support Email & Phone:** Rendered in customer app footer and contact screen.")

# Output
file_path = os.path.join(OUT_DIR, "8_User_Manual_Admin_Panel_VibeMyNight.docx")
doc.save(file_path)
print("SUCCESS: 8_User_Manual_Admin_Panel_VibeMyNight.docx generated at", file_path)
