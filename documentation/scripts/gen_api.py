import os, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_DIR = r"C:\Users\ADMIN\Downloads\vibemynight-phase1-10\documentation"
os.makedirs(OUT_DIR, exist_ok=True)

def apply_header(cell, text, fill_color="1E3A8A"):
    cell.value = text
    cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def apply_cell(cell, value, is_even=False, align="left", bold=False, color=None):
    cell.value = value
    cell.font = Font(name="Segoe UI", size=9, bold=bold, color=color or "111827")
    bg = "F9FAFB" if is_even else "FFFFFF"
    cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style="thin", color="E5E7EB"), right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"), bottom=Side(style="thin", color="E5E7EB")
    )

def fit_cols(ws):
    for col in ws.columns:
        max_l = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            for line in val_str.split("\n"):
                if len(line) > max_l:
                    max_l = len(line)
        ws.column_dimensions[col_letter].width = min(max(max_l + 3, 13), 50)

# 1. API DOCUMENTATION
wb = openpyxl.Workbook()
ws_pub = wb.active
ws_pub.title = "Public APIs"
headers_pub = ["ID", "Module", "Method", "Path", "Auth", "Params", "Request Body", "Success Response", "Error Codes", "Description"]
for i, h in enumerate(headers_pub, 1):
    apply_header(ws_pub.cell(row=1, column=i), h, "1E3A8A")
ws_pub.row_dimensions[1].height = 26

pub_rows = [
    ("PUB-001", "Events", "GET", "/api/v1/events", "None", "None", "None", "[{\"id\": 1, \"name\": \"Navratri Nights 2026\", \"slug\": \"navratri-nights-2026\", \"city\": \"Ahmedabad\", \"startDate\": \"2026-10-15\", \"endDate\": \"2026-10-19\", \"dayCount\": 5, \"startingPrice\": 499.0, \"status\": \"PUBLISHED\"}]", "500 Server Error", "Returns published events catalog for consumer homepage and festival discovery."),
    ("PUB-002", "Events", "GET", "/api/v1/events/{slug}", "None", "slug: String", "None", "{\"id\": 1, \"name\": \"Navratri Nights 2026\", \"slug\": \"navratri-nights-2026\", \"description\": \"...\", \"facilities\": [...], \"galleryImageUrls\": [...], \"highlights\": [...], \"rules\": [...]}", "404 Not Found", "Fetches complete event details including photo gallery, venue amenities, highlights and terms."),
    ("PUB-003", "Event Days", "GET", "/api/v1/events/{eventId}/days", "None", "eventId: Long", "None", "[{\"id\": 10, \"dayNumber\": 1, \"date\": \"2026-10-15\", \"programName\": \"Opening Garba Night\", \"primaryArtistName\": \"Falguni Pathak\", \"minPrice\": 499.0}]", "404 Not Found", "Returns performance day schedule and headliner artists for date selection carousel."),
    ("PUB-004", "Event Days", "GET", "/api/v1/event-days/{dayId}", "None", "dayId: Long", "None", "{\"id\": 10, \"dayNumber\": 1, \"date\": \"2026-10-15\", \"artists\": [...], \"passes\": [{\"id\": 101, \"name\": \"VIP Pass\", \"price\": 1499.0}]}", "404 Not Found", "Full details for single festival day with artists schedule and pass categories."),
    ("PUB-005", "Passes", "GET", "/api/v1/event-days/{dayId}/passes", "None", "dayId: Long", "None", "[{\"id\": 101, \"name\": \"VIP Pass\", \"price\": 1499.0, \"available\": true}]", "404 Not Found", "Available pass tiers and live pricing for the selected event day."),
    ("PUB-006", "Artists", "GET", "/api/v1/artists", "None", "None", "None", "[{\"id\": 5, \"name\": \"Falguni Pathak\", \"genre\": \"Traditional Garba\", \"imageUrl\": \"https://...\"}]", "500 Error", "Master list of performing artists and vocalists across all festival nights."),
    ("PUB-007", "Artists", "GET", "/api/v1/artists/{id}", "None", "id: Long", "None", "{\"id\": 5, \"name\": \"Falguni Pathak\", \"bio\": \"...\", \"spotifyUrl\": \"...\", \"instagramUrl\": \"...\"}", "404 Not Found", "Artist portfolio, biography, music links, and scheduled festival dates."),
    ("PUB-008", "Facilities", "GET", "/api/v1/facilities", "None", "None", "None", "[{\"id\": 1, \"name\": \"Valet Parking\", \"icon\": \"local_parking\", \"description\": \"Free valet at Gate 3\"}]", "500 Error", "Venue amenities and facility items."),
    ("PUB-009", "Settings", "GET", "/api/v1/settings/public", "None", "None", "None", "{\"brandName\": \"VibeMyNight\", \"whatsappNumber\": \"917041615131\", \"supportEmail\": \"support@vibemynight.com\"}", "500 Error", "Dynamic public branding, official WhatsApp customer service phone, and app metadata."),
    ("PUB-010", "Inquiries", "POST", "/api/v1/inquiries", "None", "None", "{\"customerName\": \"Aarav Patel\", \"customerMobile\": \"919876543210\", \"eventDayId\": 10, \"ticketCategoryId\": 101, \"quantity\": 2, \"message\": \"VIP sofa query\"}", "{\"id\": 5001, \"inquiryNumber\": \"VMN-2026-8821\", \"status\": \"NEW\", \"price\": 1499.0, \"quantity\": 2, \"estimatedTotal\": 2998.0, \"whatsappUrl\": \"https://wa.me/917041615131?text=...\"}", "400 Validation, 404 Invalid Day/Pass", "Creates customer pass inquiry. Server calculates rates from DB (anti-tamper) and returns pre-filled WhatsApp deep link."),
    ("PUB-011", "Inquiries", "GET", "/api/v1/inquiries/{inquiryNumber}", "None", "inquiryNumber: String", "None", "{\"inquiryNumber\": \"VMN-2026-8821\", \"status\": \"CONFIRMED\", \"customerName\": \"Aarav Patel\", \"estimatedTotal\": 2998.0}", "404 Not Found", "Customer inquiry lookup and status check.")
]
for r_idx, r in enumerate(pub_rows, 2):
    for c_idx, val in enumerate(r, 1):
        apply_cell(ws_pub.cell(row=r_idx, column=c_idx), val, r_idx%2==0, "center" if c_idx in [1, 3, 5] else "left", bold=(c_idx==1))
    ws_pub.row_dimensions[r_idx].height = 36
fit_cols(ws_pub)

ws_adm = wb.create_sheet(title="Admin APIs")
headers_adm = ["ID", "Module", "Method", "Path", "Auth (Role)", "Params", "Request Body", "Response (200/201)", "Errors", "Description"]
for i, h in enumerate(headers_adm, 1):
    apply_header(ws_adm.cell(row=1, column=i), h, "831843")
ws_adm.row_dimensions[1].height = 26

adm_rows = [
    ("ADM-001", "Auth", "POST", "/api/v1/auth/login", "None", "None", "{\"username\": \"admin\", \"password\": \"Admin@123\"}", "{\"accessToken\": \"eyJhbGci...\", \"tokenType\": \"Bearer\", \"name\": \"Super Admin\", \"role\": \"ROLE_ADMIN\"}", "401 Unauthorized", "Authenticates admin against BCrypt hash in DB; generates 24-hour stateless JWT token."),
    ("ADM-002", "Events", "GET", "/api/v1/admin/events", "ROLE_ADMIN", "status (optional)", "None", "[{\"id\": 1, \"name\": \"Navratri 2026\", \"status\": \"PUBLISHED\", \"inquiryCount\": 45}]", "401, 403", "Full list of events across all statuses (DRAFT, PUBLISHED, CANCELLED) with inquiry metrics."),
    ("ADM-003", "Events", "POST", "/api/v1/admin/events", "ROLE_ADMIN", "None", "{\"name\": \"Garba Fest\", \"slug\": \"garba-fest\", \"city\": \"Surat\", \"startDate\": \"2026-10-15\", \"endDate\": \"2026-10-20\"}", "{\"id\": 2, \"name\": \"Garba Fest\", \"status\": \"DRAFT\"}", "400, 409 Conflict", "Creates new festival entity in DRAFT state."),
    ("ADM-004", "Events", "PUT", "/api/v1/admin/events/{id}", "ROLE_ADMIN", "id: Long", "Full Event DTO (name, rules, highlights, gallery, venue)", "200 OK (Updated Event)", "400, 404", "Updates event metadata, gallery, terms, and highlights."),
    ("ADM-005", "Events", "PATCH", "/api/v1/admin/events/{id}/status", "ROLE_ADMIN", "id: Long", "{\"status\": \"PUBLISHED\"}", "200 OK {\"id\": 1, \"status\": \"PUBLISHED\"}", "400, 404", "Switches event publishing lifecycle. Only PUBLISHED events are visible to public."),
    ("ADM-006", "Event Days", "POST", "/api/v1/admin/events/{eventId}/days", "ROLE_ADMIN", "eventId: Long", "{\"dayNumber\": 1, \"date\": \"2026-10-15\", \"programName\": \"Inauguration Garba\"}", "201 Created (Event Day Object)", "400, 404", "Adds calendar festival day to multi-day event."),
    ("ADM-007", "Event Days", "DELETE", "/api/v1/admin/event-days/{id}", "ROLE_ADMIN", "id: Long", "None", "204 No Content", "404, 409", "Deletes festival day and cascades assigned passes and lineup."),
    ("ADM-008", "Lineup", "POST", "/api/v1/admin/event-days/{dayId}/artists", "ROLE_ADMIN", "dayId: Long", "{\"artistId\": 5, \"isPrimary\": true, \"performanceStartTime\": \"08:30 PM\"}", "200 OK (Assigned Lineup)", "404", "Assigns performer to a festival day schedule with primary headliner tag."),
    ("ADM-009", "Passes", "POST", "/api/v1/admin/event-days/{dayId}/passes", "ROLE_ADMIN", "dayId: Long", "{\"name\": \"VIP Couple Pass\", \"price\": 2499.0, \"totalCapacity\": 200}", "201 Created (Pass ID)", "400, 404", "Creates ticket tier with specific price for an event day."),
    ("ADM-010", "Passes", "PUT", "/api/v1/admin/passes/{id}", "ROLE_ADMIN", "id: Long", "{\"name\": \"VIP Early Bird\", \"price\": 1999.0, \"isActive\": true}", "200 OK (Updated Pass)", "404", "Updates ticket category pricing and live availability status."),
    ("ADM-011", "Artists", "POST", "/api/v1/admin/artists", "ROLE_ADMIN", "None", "{\"name\": \"Darshan Raval\", \"genre\": \"Bollywood / Folk\", \"imageUrl\": \"https://...\"}", "201 Created", "400", "Registers new artist in master catalog."),
    ("ADM-012", "Facilities", "POST", "/api/v1/admin/facilities", "ROLE_ADMIN", "None", "{\"name\": \"Food Court\", \"icon\": \"restaurant\", \"description\": \"20+ multi-cuisine stalls\"}", "201 Created", "400", "Adds venue amenity to master facilities list."),
    ("ADM-013", "Inquiries", "GET", "/api/v1/admin/inquiries", "ROLE_ADMIN", "status, date filters", "None", "[{\"id\": 5001, \"inquiryNumber\": \"VMN-2026-8821\", \"customerName\": \"Aarav Patel\", \"customerMobile\": \"919876543210\", \"estimatedTotal\": 2998.0, \"status\": \"NEW\"}]", "401", "Admin inquiry ledger with customer phone, pass category requested, and estimated revenue."),
    ("ADM-014", "Inquiries", "PATCH", "/api/v1/admin/inquiries/{id}/status", "ROLE_ADMIN", "id: Long", "{\"status\": \"CONFIRMED\"}", "200 OK (Updated Inquiry)", "400, 404", "Updates lead state as admin completes offline pass allocation."),
    ("ADM-015", "Settings", "PUT", "/api/v1/admin/settings", "ROLE_ADMIN", "None", "{\"brandName\": \"VibeMyNight\", \"whatsappNumber\": \"917041615131\", \"currency\": \"INR\"}", "200 OK (Updated Settings)", "400", "Updates platform configuration and booking WhatsApp phone.")
]
for r_idx, r in enumerate(adm_rows, 2):
    for c_idx, val in enumerate(r, 1):
        apply_cell(ws_adm.cell(row=r_idx, column=c_idx), val, r_idx%2==0, "center" if c_idx in [1, 3, 5] else "left", bold=(c_idx==1))
    ws_adm.row_dimensions[r_idx].height = 36
fit_cols(ws_adm)

p = os.path.join(OUT_DIR, "3_API_Documentation_VibeMyNight.xlsx")
wb.save(p)
print("SUCCESS: API Documentation Excel saved at", p)
