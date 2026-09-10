import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

OUT_DIR = r"C:\Users\ADMIN\Downloads\vibemynight-phase1-10\documentation"
os.makedirs(OUT_DIR, exist_ok=True)

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    if level == 1:
        run.font.name = 'Segoe UI'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.name = 'Segoe UI'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(131, 24, 67)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.name = 'Segoe UI'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(55, 65, 81)
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
    return h

def add_callout(doc, text, title="TIP"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F3F4F6")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="EC4899"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r_title = p.add_run(f"[{title}] ")
    r_title.bold = True
    r_title.font.name = "Segoe UI"
    r_title.font.size = Pt(9.5)
    r_title.font.color.rgb = RGBColor(131, 24, 67)
    r_text = p.add_run(text)
    r_text.font.name = "Segoe UI"
    r_text.font.size = Pt(9.5)
    r_text.font.color.rgb = RGBColor(31, 41, 55)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_table_data(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    hdr_row = tbl.rows[0]
    for idx, h in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, "1E3A8A")
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(h)
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        if col_widths and idx < len(col_widths):
            cell.width = col_widths[idx]
            
    for r_idx, row_data in enumerate(data):
        row = tbl.rows[r_idx + 1]
        bg = "F9FAFB" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(str(val))
            r.font.name = "Segoe UI"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(17, 24, 39)
            if col_widths and c_idx < len(col_widths):
                cell.width = col_widths[c_idx]
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# BUILD CUSTOMER MANUAL
doc = Document()

# Cover
title_p = doc.add_paragraph()
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
title_p.paragraph_format.space_before = Pt(36)
title_p.paragraph_format.space_after = Pt(6)
r_title = title_p.add_run("USER MANUAL — CUSTOMER WEB & MOBILE APP")
r_title.font.name = "Segoe UI"
r_title.font.size = Pt(22)
r_title.font.bold = True
r_title.font.color.rgb = RGBColor(30, 58, 138)

sub_p = doc.add_paragraph()
sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub_p.paragraph_format.space_after = Pt(24)
r_sub = sub_p.add_run("VibeMyNight — Festival Discovery, Lineup Schedules & WhatsApp Pass Booking Walkthrough")
r_sub.font.name = "Segoe UI"
r_sub.font.size = Pt(13)
r_sub.font.color.rgb = RGBColor(131, 24, 67)

meta_data = [
    ["App Version", "1.0.0 (Production)"],
    ["Access Link", "https://vibemynight.vercel.app"],
    ["Target Devices", "Desktop Web, Tablets, Android & iPhone Smartphones"],
    ["User Registration", "Not Required (Zero Login Friction)"],
    ["Booking Mode", "Direct WhatsApp Pass Allocation & Offline Payment"]
]
add_table_data(doc, ["Application Property", "Details"], meta_data, [Inches(2.2), Inches(4.3)])

doc.add_page_break()

# 1. Getting Started & App Flow
add_styled_heading(doc, "1. Customer Journey Overview", 1)
p_flow = doc.add_paragraph()
p_flow.add_run("VibeMyNight allows attendees to explore live events and secure festival passes in under 60 seconds without creating accounts or passwords:\n")

steps = [
    ("Step 1: Discover Festivals — ", "Browse featured Navratri and nightlife festivals on the home screen filtered by city (Ahmedabad, Surat, Vadodara)."),
    ("Step 2: Choose Festival & Date — ", "View festival highlights, rules, and tap on any specific performance date (e.g. Day 1: Opening Night, Day 8: Sharad Purnima)."),
    ("Step 3: Select Pass Category — ", "Compare pass tiers (VIP, Couple, Ground Access) and live prices."),
    ("Step 4: Submit Inquiry — ", "Choose quantity, enter Name and Mobile number, and tap 'Inquire via WhatsApp'."),
    ("Step 5: Connect on WhatsApp — ", "Tap the green button to open WhatsApp with a pre-filled booking summary for instant confirmation.")
]
for st, sd in steps:
    bp = doc.add_paragraph(style='List Bullet')
    r_st = bp.add_run(st)
    r_st.bold = True
    r_st.font.color.rgb = RGBColor(30, 58, 138)
    bp.add_run(sd)

# 2. Screen-by-Screen Walkthrough
add_styled_heading(doc, "2. Screen-by-Screen Walkthrough", 1)

add_styled_heading(doc, "2.1 Homepage (Route: /)", 2)
p_home = doc.add_paragraph()
p_home.add_run("The homepage features high-energy dark neon aesthetics with smooth transitions:")
h_points = [
    ("Hero Featured Carousel: ", "Displays top headliner festivals with direct 'Book Passes' and 'Explore Lineup' action buttons."),
    ("Upcoming Nights Grid: ", "Chronological festival cards showing starting prices, city badges, and date counts."),
    ("Featured Artists Spotlight: ", "Horizontal list of popular singers and DJs with genre tags."),
    ("Dynamic Navigation Navbar: ", "Instant navigation to Events, Artists, About, and Contact pages.")
]
for ht, hd in h_points:
    bp = doc.add_paragraph(style='List Bullet')
    r_ht = bp.add_run(ht)
    r_ht.bold = True
    bp.add_run(hd)

add_styled_heading(doc, "2.2 Festival Details Screen (Route: /events/:slug)", 2)
p_evt = doc.add_paragraph()
p_evt.add_run("Deep details page for an individual festival:")
e_points = [
    ("Photo Gallery Carousel: ", "High-res images showcasing arena decor, stage lighting, and previous year celebrations."),
    ("Venue Amenities & Facilities: ", "Icons for Free Valet Parking, Air-Conditioned VIP Lounges, Food Court, and Security."),
    ("Event Highlights & Guidelines: ", "Bullet points regarding dress codes, traditional attire, and gate timings."),
    ("Calendar Date Selector: ", "Interactive horizontal date cards displaying daily starting prices and headliner vocalists.")
]
for et, ed in e_points:
    bp = doc.add_paragraph(style='List Bullet')
    r_et = bp.add_run(et)
    r_et.bold = True
    bp.add_run(ed)

add_styled_heading(doc, "2.3 Single Day & Pass Selection (Route: /event-days/:dayId)", 2)
p_day = doc.add_paragraph()
p_day.add_run("Dedicated screen for a selected festival performance day:")
d_points = [
    ("Headliner & Lineup: ", "Displays performance order, vocalists, and stage time slots."),
    ("Pass Category Cards: ", "VIP, Couple, and General access cards with included benefits and clear INR prices."),
    ("Direct 'Get Passes' Action: ", "Instantly takes the user into the booking inquiry flow with pre-selected pass tier.")
]
for dt, dd in d_points:
    bp = doc.add_paragraph(style='List Bullet')
    r_dt = bp.add_run(dt)
    r_dt.bold = True
    bp.add_run(dd)

add_styled_heading(doc, "2.4 Pass Inquiry Screen (Route: /inquiry)", 2)
p_inq = doc.add_paragraph()
p_inq.add_run("The frictionless booking submission form:")
i_points = [
    ("Live Quantity Stepper: ", "Interactive (-) and (+) buttons adjusting ticket quantity with live price updates."),
    ("Customer Information Fields: ", "Name (required), WhatsApp Mobile Number (required), Email (optional), Custom notes/requests."),
    ("Server Rate Verification: ", "Transparent calculation breakdown showing Price x Quantity = Total."),
    ("Submit Action: ", "Securely submits inquiry to Spring Boot backend.")
]
for it, idesc in i_points:
    bp = doc.add_paragraph(style='List Bullet')
    r_it = bp.add_run(it)
    r_it.bold = True
    bp.add_run(idesc)

add_styled_heading(doc, "2.5 Confirmation & WhatsApp Chat (Route: /inquiry/success)", 2)
p_suc = doc.add_paragraph()
p_suc.add_run("Post-submission confirmation screen:")
s_points = [
    ("Inquiry Reference Card: ", "Displays unique booking ID (e.g. `VMN-2026-8821`) and confirmed order summary."),
    ("One-Tap 'Continue to WhatsApp' Button: ", "Opens WhatsApp Web or mobile app with pre-filled message addressed to official organizer helpline (`917041615131`)."),
    ("Direct Helpline Call: ", "Provides alternative direct phone support option.")
]
for st2, sd2 in s_points:
    bp = doc.add_paragraph(style='List Bullet')
    r_st2 = bp.add_run(st2)
    r_st2.bold = True
    bp.add_run(sd2)

add_callout(doc, "If WhatsApp Web is blocked on corporate networks, customers can directly quote their Inquiry Number (e.g. VMN-2026-8821) via SMS or phone call to complete pass pickup.", "OFFLINE FALLBACK")

# Output
file_path = os.path.join(OUT_DIR, "7_User_Manual_Customer_App_VibeMyNight.docx")
doc.save(file_path)
print("SUCCESS: 7_User_Manual_Customer_App_VibeMyNight.docx generated at", file_path)
