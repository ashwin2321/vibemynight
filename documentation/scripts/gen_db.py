import os, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT_DIR = r"C:\Users\ADMIN\Downloads\vibemynight-phase1-10\documentation"
os.makedirs(OUT_DIR, exist_ok=True)

def apply_header(cell, text, fill_color="065F46"):
    cell.value = text
    cell.font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def apply_cell(cell, value, is_even=False, align="left", bold=False):
    cell.value = value
    cell.font = Font(name="Segoe UI", size=9, bold=bold, color="111827")
    bg = "F9FAFB" if is_even else "FFFFFF"
    cell.fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = Border(
        left=Side(style="thin", color="E5E7EB"), right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"), bottom=Side(style="thin", color="E5E7EB")
    )

def fit_cols(ws):
    for col in ws.columns:
        max_l = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or "")
            for line in val_str.split("\n"):
                if len(line) > max_l: max_l = len(line)
        ws.column_dimensions[col_letter].width = min(max(max_l + 3, 13), 50)

wb = openpyxl.Workbook()
ws_rel = wb.active
ws_rel.title = "ERD & Relationships"
headers_rel = ["ID", "Parent Table", "PK Column", "Child Table", "FK Column", "Cardinality", "Cascade", "Business Rule & Usage"]
for i, h in enumerate(headers_rel, 1):
    apply_header(ws_rel.cell(row=1, column=i), h, "065F46")
ws_rel.row_dimensions[1].height = 26

rels = [
    ("REL-01", "events", "id", "event_days", "event_id", "1:N", "CASCADE", "Festival consists of multiple calendar performance days."),
    ("REL-02", "event_days", "id", "ticket_categories", "event_day_id", "1:N", "CASCADE", "Each specific date has distinct pass categories and pricing."),
    ("REL-03", "event_days", "id", "event_day_artists", "event_day_id", "1:N", "CASCADE", "Junction table linking performing artists to festival days."),
    ("REL-04", "artists", "id", "event_day_artists", "artist_id", "1:N", "RESTRICT", "Master artist catalog profile mapped to schedule."),
    ("REL-05", "events", "id", "event_facilities", "event_id", "1:N", "CASCADE", "Venue amenities offered at the festival venue."),
    ("REL-06", "facilities", "id", "event_facilities", "facility_id", "1:N", "RESTRICT", "Master facility catalog reference."),
    ("REL-07", "events", "id", "inquiries", "event_id", "1:N", "SET NULL", "Maintains customer booking history against event."),
    ("REL-08", "event_days", "id", "inquiries", "event_day_id", "1:N", "RESTRICT", "Customer selected festival date."),
    ("REL-09", "ticket_categories", "id", "inquiries", "ticket_category_id", "1:N", "RESTRICT", "Pass category (VIP/General) requested by customer."),
    ("REL-10", "events", "id", "event_gallery_images", "event_id", "1:N", "CASCADE", "Event photo gallery URLs."),
    ("REL-11", "events", "id", "event_highlights", "event_id", "1:N", "CASCADE", "Festival bullet highlights and attractions."),
    ("REL-12", "events", "id", "event_rules", "event_id", "1:N", "CASCADE", "Terms of entry, dress codes, and security rules.")
]
for r_idx, row in enumerate(rels, 2):
    for c_idx, val in enumerate(row, 1):
        apply_cell(ws_rel.cell(row=r_idx, column=c_idx), val, r_idx%2==0, "center" if c_idx in [1, 3, 5, 6, 7] else "left", bold=(c_idx==1))
    ws_rel.row_dimensions[r_idx].height = 25
fit_cols(ws_rel)

tables = {
    "events": [
        ("id", "BIGINT AUTO_INCREMENT", "PRIMARY KEY", "Unique festival event ID"),
        ("name", "VARCHAR(255)", "NOT NULL", "Event title (e.g. Navratri Nights 2026)"),
        ("slug", "VARCHAR(255)", "NOT NULL, UNIQUE", "SEO friendly URL slug identifier"),
        ("description", "TEXT", "NULL", "Comprehensive festival description and theme"),
        ("venue", "VARCHAR(255)", "NOT NULL", "Venue ground / stadium name"),
        ("location", "VARCHAR(255)", "NULL", "Street address / landmark"),
        ("city", "VARCHAR(100)", "NOT NULL", "City location (Ahmedabad, Surat, etc.)"),
        ("start_date", "DATE", "NOT NULL", "Opening date of festival"),
        ("end_date", "DATE", "NOT NULL", "Closing date of festival"),
        ("main_image", "VARCHAR(500)", "NULL", "Hero banner image URL"),
        ("thumbnail", "VARCHAR(500)", "NULL", "Card thumbnail image URL"),
        ("featured", "BOOLEAN", "DEFAULT FALSE", "Homepage hero featured status"),
        ("status", "VARCHAR(50)", "DEFAULT \"DRAFT\"", "Lifecycle: DRAFT, PUBLISHED, CANCELLED, ARCHIVED"),
        ("created_at", "DATETIME", "DEFAULT CURRENT_TIMESTAMP", "Audit creation timestamp"),
        ("updated_at", "DATETIME", "ON UPDATE CURRENT_TIMESTAMP", "Audit update timestamp")
    ],
    "event_days": [
        ("id", "BIGINT AUTO_INCREMENT", "PRIMARY KEY", "Event day identifier"),
        ("event_id", "BIGINT", "FOREIGN KEY -> events(id)", "Reference to parent festival event"),
        ("day_number", "INT", "NOT NULL", "Sequential day index (1 to N)"),
        ("date", "DATE", "NOT NULL", "Calendar performance date"),
        ("program_name", "VARCHAR(255)", "NULL", "Special day theme / program name"),
        ("start_time", "VARCHAR(50)", "NULL", "Gate open / start time (07:30 PM)"),
        ("end_time", "VARCHAR(50)", "NULL", "Curfew / end time (11:45 PM)"),
        ("status", "VARCHAR(50)", "DEFAULT \"ACTIVE\"", "ACTIVE, SOLD_OUT, CANCELLED"),
        ("created_at", "DATETIME", "DEFAULT CURRENT_TIMESTAMP", "Audit timestamp")
    ],
    "ticket_categories": [
        ("id", "BIGINT AUTO_INCREMENT", "PRIMARY KEY", "Pass category ID"),
        ("event_day_id", "BIGINT", "FOREIGN KEY -> event_days(id)", "Linked performance date"),
        ("name", "VARCHAR(150)", "NOT NULL", "Pass tier (VIP, General, Couple, Student)"),
        ("description", "VARCHAR(500)", "NULL", "Perks and inclusions description"),
        ("price", "DECIMAL(10,2)", "NOT NULL, CHECK (price >= 0)", "Base price per pass in INR"),
        ("total_capacity", "INT", "NULL", "Pass allocation quota"),
        ("sold_count", "INT", "DEFAULT 0", "Count of confirmed tickets"),
        ("is_active", "BOOLEAN", "DEFAULT TRUE", "Toggles pass category availability"),
        ("created_at", "DATETIME", "DEFAULT CURRENT_TIMESTAMP", "Audit timestamp")
    ],
    "artists": [
        ("id", "BIGINT AUTO_INCREMENT", "PRIMARY KEY", "Artist record ID"),
        ("name", "VARCHAR(150)", "NOT NULL", "Artist / Performer / Band stage name"),
        ("genre", "VARCHAR(100)", "NULL", "Genre (Traditional Garba, Folk, Bollywood)"),
        ("bio", "TEXT", "NULL", "Artist profile and biography"),
        ("image_url", "VARCHAR(500)", "NULL", "High-res profile photo URL"),
        ("spotify_url", "VARCHAR(500)", "NULL", "Spotify profile URL"),
        ("instagram_url", "VARCHAR(500)", "NULL", "Instagram profile URL"),
        ("status", "VARCHAR(50)", "DEFAULT \"ACTIVE\"", "ACTIVE / INACTIVE flag"),
        ("created_at", "DATETIME", "DEFAULT CURRENT_TIMESTAMP", "Audit timestamp")
    ],
    "inquiries": [
        ("id", "BIGINT AUTO_INCREMENT", "PRIMARY KEY", "Inquiry lead ID"),
        ("inquiry_number", "VARCHAR(64)", "NOT NULL, UNIQUE", "Customer reference code (VMN-2026-8821)"),
        ("customer_name", "VARCHAR(150)", "NOT NULL", "Customer full name"),
        ("customer_mobile", "VARCHAR(20)", "NOT NULL", "Customer WhatsApp number"),
        ("customer_email", "VARCHAR(150)", "NULL", "Customer email address"),
        ("event_id", "BIGINT", "FOREIGN KEY -> events(id)", "Target festival"),
        ("event_day_id", "BIGINT", "FOREIGN KEY -> event_days(id)", "Target festival day"),
        ("ticket_category_id", "BIGINT", "FOREIGN KEY -> ticket_categories(id)", "Target pass tier"),
        ("unit_price", "DECIMAL(10,2)", "NOT NULL", "Server-computed rate at moment of inquiry"),
        ("quantity", "INT", "NOT NULL, CHECK (quantity > 0)", "Pass quantity requested"),
        ("estimated_total", "DECIMAL(10,2)", "NOT NULL", "Server-calculated total (unit_price * quantity)"),
        ("customer_message", "TEXT", "NULL", "Custom inquiry notes"),
        ("status", "VARCHAR(50)", "DEFAULT \"NEW\"", "NEW, CONTACTED, CONFIRMED, CANCELLED"),
        ("whatsapp_url", "VARCHAR(1000)", "NULL", "Pre-rendered direct WhatsApp chat link"),
        ("created_at", "DATETIME", "DEFAULT CURRENT_TIMESTAMP", "Audit submission timestamp")
    ],
    "app_settings": [
        ("id", "BIGINT AUTO_INCREMENT", "PRIMARY KEY", "Singleton settings row"),
        ("brand_name", "VARCHAR(100)", "DEFAULT \"VibeMyNight\"", "Public platform brand name"),
        ("whatsapp_number", "VARCHAR(20)", "NOT NULL", "Official helpline WhatsApp number (917041615131)"),
        ("support_email", "VARCHAR(100)", "NULL", "Customer care support email"),
        ("support_phone", "VARCHAR(50)", "NULL", "Direct phone call helpline"),
        ("currency", "VARCHAR(10)", "DEFAULT \"INR\"", "Currency ISO code"),
        ("updated_at", "DATETIME", "ON UPDATE CURRENT_TIMESTAMP", "Audit timestamp")
    ],
    "admin_users": [
        ("id", "BIGINT AUTO_INCREMENT", "PRIMARY KEY", "Admin user ID"),
        ("username", "VARCHAR(100)", "NOT NULL, UNIQUE", "Login username"),
        ("password_hash", "VARCHAR(255)", "NOT NULL", "BCrypt hashed password (12 rounds)"),
        ("full_name", "VARCHAR(150)", "NOT NULL", "Administrator full name"),
        ("email", "VARCHAR(150)", "NOT NULL, UNIQUE", "Admin email"),
        ("role", "VARCHAR(50)", "DEFAULT \"ROLE_ADMIN\"", "RBAC role"),
        ("is_active", "BOOLEAN", "DEFAULT TRUE", "Account active flag")
    ]
}

for t_name, cols in tables.items():
    ws = wb.create_sheet(title=t_name)
    h_list = ["Column Name", "Data Type", "Constraints & Keys", "Business Description & Usage"]
    for i, h in enumerate(h_list, 1):
        apply_header(ws.cell(row=1, column=i), h, "1E293B")
    ws.row_dimensions[1].height = 26
    for r_idx, c_info in enumerate(cols, 2):
        for c_idx, val in enumerate(c_info, 1):
            apply_cell(ws.cell(row=r_idx, column=c_idx), val, r_idx%2==0, "center" if c_idx in [2, 3] else "left", bold=(c_idx==1))
        ws.row_dimensions[r_idx].height = 24
    fit_cols(ws)

p = os.path.join(OUT_DIR, "4_Database_Schema_Documentation_VibeMyNight.xlsx")
wb.save(p)
print("SUCCESS: Database Schema Documentation Excel saved at", p)
